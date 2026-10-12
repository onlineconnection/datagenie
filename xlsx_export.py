import io, datetime, decimal
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter
import db

NAVY = "0B1F44"
BRAND = "1D4ED8"
MONEY_WORDS = ("price", "amount", "cost", "total", "salary", "fee", "balance", "revenue", "paid",
               "payment", "wage", "rent", "budget", "sales", "income", "expense", "profit",
               "charge", "fare", "tuition", "deposit", "loan", "interest")


def pretty(n):
    return n.replace("_", " ").title()


def ismoney(name):
    return any(w in name for w in MONEY_WORDS)


def build(biz, symbol, code):
    money_fmt = '"%s"#,##0.00' % symbol
    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True
    ov = wb.active
    ov.title = "Overview"
    ov.sheet_properties.tabColor = NAVY
    tables = biz["design"]["tables"]
    sheets = {}
    for t in tables:
        ws = wb.create_sheet(pretty(t["name"])[:31])
        ws.sheet_properties.tabColor = BRAND
        sheets[t["name"]] = ws
    labelidx = {}
    for t in tables:
        names = [c["name"] for c in t["columns"]]
        try:
            lc = db.label_col(t)
        except Exception:
            lc = names[0]
        labelidx[t["name"]] = names.index(lc) if lc in names else 0

    for t in tables:
        ws = sheets[t["name"]]
        cols = t["columns"]
        rows = db.search_rows(biz, t, "", limit=50000)
        opts = db.fk_options(biz, t)
        widths = []
        for j, c in enumerate(cols, 1):
            h = "ID" if c["primary_key"] else pretty(c["name"])
            ws.cell(row=1, column=j, value=h)
            widths.append(len(h))
        ws.row_dimensions[1].height = 22
        for i, r in enumerate(rows, 2):
            for j, c in enumerate(cols, 1):
                v = r[j - 1]
                if v is None:
                    continue
                if c["type"] == "boolean":
                    v = "Yes" if v else "No"
                elif c["references"]:
                    v = opts.get(c["name"], {}).get(v, v)
                elif isinstance(v, decimal.Decimal):
                    v = float(v)
                cell = ws.cell(row=i, column=j)
                cell.value = v
                if isinstance(v, str) and v.startswith("="):
                    cell.data_type = "s"
                if c["type"] == "numeric":
                    cell.number_format = money_fmt if ismoney(c["name"]) else "#,##0.00"
                elif c["type"] == "date":
                    cell.number_format = "yyyy-mm-dd"
                elif c["type"] == "timestamp":
                    cell.number_format = "yyyy-mm-dd hh:mm"
                widths[j - 1] = max(widths[j - 1], len(str(v)))
        last = max(len(rows) + 1, 2)
        tb = Table(displayName="t_" + t["name"], ref="A1:%s%d" % (get_column_letter(len(cols)), last))
        tb.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tb)
        ws.freeze_panes = "A2"
        for j, c in enumerate(cols, 1):
            L = get_column_letter(j)
            ws.column_dimensions[L].width = min(max(widths[j - 1] + 3, 12), 45)
            dv = None
            if c["type"] == "boolean":
                dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
            elif c["references"]:
                parent = c["references"].split(".")[0]
                if parent in sheets:
                    pl = get_column_letter(labelidx[parent] + 1)
                    dv = DataValidation(type="list", allow_blank=True,
                                        formula1="='%s'!$%s$2:$%s$5000" % (sheets[parent].title, pl, pl))
            if dv:
                ws.add_data_validation(dv)
                dv.add("%s2:%s1000" % (L, L))

    ov["A1"] = biz["title"]
    ov["A1"].font = Font(size=22, bold=True, color=NAVY)
    ov["A2"] = "Exported from DataGenie on %s  |  Currency: %s" % (datetime.date.today().isoformat(), code)
    ov["A2"].font = Font(size=11, color="52627F")
    head_fill = PatternFill("solid", fgColor=BRAND)
    head_font = Font(bold=True, color="FFFFFF")
    ov["A4"], ov["B4"] = "Table", "Records"
    for a in ("A4", "B4"):
        ov[a].fill, ov[a].font = head_fill, head_font
    r = 5
    for t in tables:
        ws = sheets[t["name"]]
        ov.cell(row=r, column=1, value=pretty(t["name"]))
        ov.cell(row=r, column=1).hyperlink = "#'%s'!A1" % ws.title
        ov.cell(row=r, column=1).font = Font(bold=True, color=BRAND, underline="single")
        ov.cell(row=r, column=2, value="=COUNT('%s'!$A$2:$A$5000)" % ws.title)
        r += 1
    n = len(tables)
    money_rows = []
    for t in tables:
        ws = sheets[t["name"]]
        for j, c in enumerate(t["columns"], 1):
            if c["type"] == "numeric" and ismoney(c["name"]) and len(money_rows) < 12:
                L = get_column_letter(j)
                money_rows.append(("%s: %s" % (pretty(t["name"]), pretty(c["name"])),
                                   "=SUM('%s'!$%s$2:$%s$5000)" % (ws.title, L, L)))
    if money_rows:
        r += 1
        ov.cell(row=r, column=1, value="Money totals")
        ov.cell(row=r, column=2, value="Amount")
        for col in (1, 2):
            ov.cell(row=r, column=col).fill, ov.cell(row=r, column=col).font = head_fill, head_font
        for label, f in money_rows:
            r += 1
            ov.cell(row=r, column=1, value=label)
            c = ov.cell(row=r, column=2, value=f)
            c.number_format = money_fmt
    ov.column_dimensions["A"].width = 36
    ov.column_dimensions["B"].width = 18
    ch = BarChart()
    ch.type = "bar"
    ch.title = "Records per table"
    ch.legend = None
    ch.add_data(Reference(ov, min_col=2, min_row=4, max_row=4 + n), titles_from_data=True)
    ch.set_categories(Reference(ov, min_col=1, min_row=5, max_row=4 + n))
    ch.height, ch.width = 8, 16
    ov.add_chart(ch, "D4")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
