import csv, io, re, decimal, datetime
from openpyxl import load_workbook

MAX_SHEETS = 15
MAX_COLS = 25
MAX_ROWS = 2000
MAX_TOTAL = 10000
EMPTY = (None, "")


def clean(name):
    s = re.sub(r"[^a-z0-9]+", "_", str(name).strip().lower()).strip("_")
    if s and not s[0].isalpha():
        s = "f_" + s
    return s[:40]


def read_sheets(data, filename):
    out = []
    if filename.lower().endswith(".csv"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("latin-1")
        try:
            delim = csv.Sniffer().sniff(text[:4096], delimiters=",;\t").delimiter
        except csv.Error:
            delim = ","
        rows = [tuple(r) for r in csv.reader(io.StringIO(text), delimiter=delim)][:MAX_ROWS + 11]
        out.append((filename.rsplit(".", 1)[0], rows))
        return out
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    for ws in wb.worksheets[:MAX_SHEETS]:
        if ws.sheet_state != "visible":
            continue
        rows = [tuple(r) for r in ws.iter_rows(values_only=True, max_row=MAX_ROWS + 11)]
        out.append((ws.title, rows))
    return out


def split_header(rows):
    head = rows[:10]
    counts = [sum(1 for v in r if v not in EMPTY) for r in head]
    if not counts or max(counts) == 0:
        return None
    i = counts.index(max(counts))
    return rows[i], rows[i + 1:]


def asnum(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float, decimal.Decimal)):
        f = float(v)
        return f if abs(f) < 1e11 else None
    if isinstance(v, str):
        t = v.strip().replace(",", "").replace("$", "")
        if re.fullmatch(r"-?\d+(\.\d+)?", t) and not re.match(r"-?0\d", t) and len(t) <= 15:
            f = float(t)
            return f if abs(f) < 1e11 else None
    return None


def isbool(v):
    return isinstance(v, bool) or (isinstance(v, str) and v.strip().lower() in ("yes", "no", "true", "false"))


def infer(values):
    vals = [v for v in values if v not in EMPTY]
    if not vals:
        return "text"
    if all(isbool(v) for v in vals):
        return "boolean"
    if all(isinstance(v, (datetime.datetime, datetime.date)) for v in vals):
        plain = all(not isinstance(v, datetime.datetime) or
                    (v.hour == 0 and v.minute == 0 and v.second == 0) for v in vals)
        return "date" if plain else "timestamp"
    nums = [asnum(v) for v in vals]
    if all(n is not None for n in nums):
        if all(n == int(n) and abs(n) < 2e9 for n in nums):
            return "integer"
        return "numeric"
    return "text"


def convert(v, typ):
    if v in EMPTY:
        return None
    if typ == "boolean":
        return v if isinstance(v, bool) else str(v).strip().lower() in ("yes", "true")
    if typ == "integer":
        return int(asnum(v))
    if typ == "numeric":
        return round(asnum(v), 2)
    if typ == "date":
        return v.date().isoformat() if isinstance(v, datetime.datetime) else v.isoformat()
    if typ == "timestamp":
        return v.isoformat(sep=" ")
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.isoformat()
    return str(v).strip()[:2000]


def unique(name, used):
    base, k = name, 2
    while name in used:
        name = "%s_%d" % (base[:36], k)
        k += 1
    used.add(name)
    return name


def build(data, filename):
    tables, rowsmap, tnames, total = [], {}, set(), 0
    for sname, rows in read_sheets(data, filename):
        if any(r and isinstance(r[0], str) and r[0].startswith("Exported from DataGenie") for r in rows[:3]):
            continue
        picked = split_header(rows)
        if not picked:
            continue
        head, body = picked
        n = len(head)
        while n and head[n - 1] in EMPTY:
            n -= 1
        n = min(n, MAX_COLS)
        if n == 0:
            continue
        used = {"id"}
        names = []
        for j in range(n):
            nm = clean(head[j]) if head[j] not in EMPTY else ""
            if not nm:
                nm = "column_%d" % (j + 1)
            if nm == "id":
                nm = "ref_id"
            names.append(unique(nm, used))
        grid = []
        for r in body[:MAX_ROWS]:
            row = [(r[j] if j < len(r) else None) for j in range(n)]
            if any(v not in EMPTY for v in row):
                grid.append(row)
        if not grid:
            continue
        types = [infer([row[j] for row in grid]) for j in range(n)]
        tname = unique(clean(sname) or "table_%d" % (len(tables) + 1), tnames)
        cols = [{"name": "id", "type": "integer", "required": True, "primary_key": True, "references": None}]
        cols += [{"name": names[j], "type": types[j], "required": False, "primary_key": False,
                  "references": None} for j in range(n)]
        tables.append({"name": tname, "columns": cols})
        rowsmap[tname] = [{names[j]: convert(row[j], types[j]) for j in range(n)} for row in grid]
        total += len(grid)
        if total > MAX_TOTAL:
            raise ValueError("This file has too many rows. The limit is %d in total." % MAX_TOTAL)
    if not tables:
        raise ValueError("No data found. Put column headers in the first row and records below them.")
    title = clean(filename.rsplit(".", 1)[0]) or "imported_data"
    return link({"name": title, "tables": tables}, rowsmap)


def stem(n):
    n = n.lower()
    if n.endswith("ies"):
        return n[:-3] + "y"
    if n.endswith("s") and len(n) > 3:
        return n[:-1]
    return n


def norm(v):
    return " ".join(str(v).lower().split())


def link(schema, rowsmap):
    tables = {t["name"]: t for t in schema["tables"]}
    refs = {n: set() for n in tables}

    def reaches(a, b, seen=None):
        seen = seen or set()
        if a == b:
            return True
        seen.add(a)
        return any(reaches(x, b, seen) for x in refs[a] if x not in seen)

    labels = {}
    for pn, pt in tables.items():
        lc = next((c["name"] for c in pt["columns"] if c["type"] == "text" and not c["primary_key"]), None)
        if not lc:
            continue
        vals = [norm(r.get(lc)) for r in rowsmap[pn] if r.get(lc) not in EMPTY]
        if len(vals) == len(rowsmap[pn]) and len(set(vals)) == len(vals) and vals:
            labels[pn] = {v: i + 1 for i, v in enumerate(norm(r.get(lc)) for r in rowsmap[pn])}
    for tn, t in tables.items():
        for c in t["columns"]:
            if c["type"] != "text" or c["primary_key"] or c["references"]:
                continue
            vals = [r.get(c["name"]) for r in rowsmap[tn]]
            filled = [norm(v) for v in vals if v not in EMPTY]
            if not filled:
                continue
            for pn, lab in labels.items():
                sp = stem(pn)
                nm = c["name"]
                if pn == tn or not (sp in nm or (nm in sp and len(nm) >= 4)):
                    continue
                if not all(v in lab for v in filled) or reaches(pn, tn):
                    continue
                for r in rowsmap[tn]:
                    v = r.get(nm)
                    r[nm] = lab[norm(v)] if v not in EMPTY else None
                c["type"] = "integer"
                c["references"] = pn + ".id"
                refs[tn].add(pn)
                break
    return schema, rowsmap
